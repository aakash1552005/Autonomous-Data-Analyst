"""
agents/report/pptx_generator.py
===============================
Deterministic, executive PowerPoint (PPTX) presentation generator.
Transforms verified Dataset Intelligence Object (DIO) data into a
16:9 widescreen business presentation using python-pptx.
"""

from __future__ import annotations

from pathlib import Path
from typing import Any

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

from agents.report.data_extractor import extract_report_payload

# Corporate Color Palette
COLOR_NAVY = RGBColor(0x1A, 0x36, 0x5D)      # Primary Headers
COLOR_BLUE = RGBColor(0x2B, 0x6C, 0xB0)      # Accents & Subsections
COLOR_SLATE = RGBColor(0x2D, 0x37, 0x48)     # Body Text
COLOR_MUTED = RGBColor(0x71, 0x80, 0x96)     # Labels & Footers
COLOR_CARD_BG = RGBColor(0xF7, 0xFA, 0xFC)   # Containers
COLOR_BORDER = RGBColor(0xCB, 0xD5, 0xE0)    # Card Outlines
COLOR_WHITE = RGBColor(0xFF, 0xFF, 0xFF)
COLOR_ALERT = RGBColor(0xC5, 0x30, 0x30)     # Badges & PII indicators


def _add_slide_header(slide: Any, title: str, category: str = "AUTONOMOUS DATA ANALYST") -> None:
    """Add a standardized top header banner to a presentation slide."""
    header_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(1.1))
    tf = header_box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

    # Category Eyebrow
    p_cat = tf.paragraphs[0]
    p_cat.text = category.upper()
    p_cat.font.size = Pt(9)
    p_cat.font.bold = True
    p_cat.font.color.rgb = COLOR_BLUE

    # Slide Title
    p_title = tf.add_paragraph()
    p_title.text = title
    p_title.font.size = Pt(20)
    p_title.font.bold = True
    p_title.font.color.rgb = COLOR_NAVY


def _add_footer(slide: Any, current_slide: int, total_slides: int, file_name: str) -> None:
    """Add standard bottom metadata footer to slide."""
    footer_box = slide.shapes.add_textbox(Inches(0.8), Inches(6.9), Inches(11.733), Inches(0.4))
    tf = footer_box.text_frame
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.text = f"Dataset: {file_name}   |   Confidential Analytical Intelligence   |   Slide {current_slide} of {total_slides}"
    p.font.size = Pt(8)
    p.font.color.rgb = COLOR_MUTED


def _add_kpi_card(slide: Any, left: Inches, top: Inches, width: Inches, height: Inches, value: str, label: str) -> Any:
    """Draw a rounded rectangle KPI metric card."""
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    card.fill.solid()
    card.fill.fore_color.rgb = COLOR_CARD_BG
    card.line.color.rgb = COLOR_BORDER
    card.line.width = Pt(0.75)

    tf = card.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.1)
    tf.margin_top = Inches(0.15)
    tf.margin_bottom = Inches(0.1)

    p_val = tf.paragraphs[0]
    p_val.text = value
    p_val.alignment = PP_ALIGN.CENTER
    p_val.font.size = Pt(20)
    p_val.font.bold = True
    p_val.font.color.rgb = COLOR_NAVY

    p_lbl = tf.add_paragraph()
    p_lbl.text = label
    p_lbl.alignment = PP_ALIGN.CENTER
    p_lbl.font.size = Pt(8.5)
    p_lbl.font.color.rgb = COLOR_MUTED
    return card


def _add_content_card(slide: Any, left: Inches, top: Inches, width: Inches, height: Inches, title: str, bullets: list[str]) -> Any:
    """Draw a styled content container with bulleted points."""
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    card.fill.solid()
    card.fill.fore_color.rgb = COLOR_CARD_BG
    card.line.color.rgb = COLOR_BORDER
    card.line.width = Pt(0.75)

    tf = card.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.2)
    tf.margin_top = tf.margin_bottom = Inches(0.15)

    if title:
        p_t = tf.paragraphs[0]
        p_t.text = title
        p_t.font.size = Pt(12)
        p_t.font.bold = True
        p_t.font.color.rgb = COLOR_NAVY
        p_t.space_after = Pt(6)

    for i, b in enumerate(bullets):
        p = tf.add_paragraph() if (title or i > 0) else tf.paragraphs[0]
        p.text = f"•  {b}"
        p.font.size = Pt(9.5)
        p.font.color.rgb = COLOR_SLATE
        p.space_after = Pt(4)
    return card


def _add_table(
    slide: Any,
    left: Inches,
    top: Inches,
    width: Inches,
    height: Inches,
    headers: list[str],
    rows: list[list[str]],
    col_widths: list[Inches] | None = None,
) -> Any:
    """Add a cleanly styled data table to the slide."""
    num_rows = len(rows) + 1
    num_cols = len(headers)
    table_shape = slide.shapes.add_table(num_rows, num_cols, left, top, width, height)
    tbl = table_shape.table

    if col_widths and len(col_widths) == num_cols:
        for idx, cw in enumerate(col_widths):
            tbl.columns[idx].width = cw

    # Header Row
    for c_idx, h_text in enumerate(headers):
        cell = tbl.cell(0, c_idx)
        cell.fill.solid()
        cell.fill.fore_color.rgb = COLOR_NAVY
        tf = cell.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = Inches(0.08)
        tf.margin_top = tf.margin_bottom = Inches(0.05)
        p = tf.paragraphs[0]
        p.text = h_text
        p.font.bold = True
        p.font.size = Pt(9)
        p.font.color.rgb = COLOR_WHITE

    # Data Rows
    for r_idx, r_data in enumerate(rows, start=1):
        bg_color = COLOR_WHITE if r_idx % 2 == 1 else COLOR_CARD_BG
        for c_idx, val in enumerate(r_data):
            cell = tbl.cell(r_idx, c_idx)
            cell.fill.solid()
            cell.fill.fore_color.rgb = bg_color
            tf = cell.text_frame
            tf.word_wrap = True
            tf.margin_left = tf.margin_right = Inches(0.08)
            tf.margin_top = tf.margin_bottom = Inches(0.05)
            p = tf.paragraphs[0]
            p.text = str(val)
            p.font.size = Pt(8)
            if "[Protected" in str(val):
                p.font.color.rgb = COLOR_ALERT
                p.font.bold = True
            else:
                p.font.color.rgb = COLOR_SLATE

    return table_shape


def generate_pptx_report(dio: Any, output_path: str | Path) -> Path:
    """
    Generate an executive 16:9 PowerPoint deck from DIO data.
    
    Parameters:
        dio: Dataset Intelligence Object (dataclass or dict)
        output_path: Path where the generated PPTX will be saved
        
    Returns:
        Path to the saved PPTX presentation.
    """
    dest_path = Path(output_path).resolve()
    dest_path.parent.mkdir(parents=True, exist_ok=True)

    data = extract_report_payload(dio)
    file_name = data["metadata"]["file_name"]
    dataset_id = data["metadata"]["dataset_id"]
    gen_time = data["metadata"]["generated_at"]

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Target: 11 Slides
    TOTAL_SLIDES = 11

    # =========================================================================
    # SLIDE 1: TITLE SLIDE
    # =========================================================================
    s1 = prs.slides.add_slide(blank_layout)
    # Background accent banner
    banner = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
    banner.fill.solid()
    banner.fill.fore_color.rgb = COLOR_NAVY
    banner.line.fill.background()

    # Title Box
    t_box = s1.shapes.add_textbox(Inches(1.2), Inches(2.0), Inches(10.9), Inches(3.5))
    tf1 = t_box.text_frame
    tf1.word_wrap = True

    p_eyebrow = tf1.paragraphs[0]
    p_eyebrow.text = "AUTONOMOUS DATA ANALYST — EXECUTIVE DELIVERABLE"
    p_eyebrow.font.size = Pt(13)
    p_eyebrow.font.bold = True
    p_eyebrow.font.color.rgb = RGBColor(0x63, 0xB3, 0xED)
    p_eyebrow.space_after = Pt(12)

    p_maintitle = tf1.add_paragraph()
    p_maintitle.text = f"Executive Intelligence Deck: {file_name}"
    p_maintitle.font.size = Pt(32)
    p_maintitle.font.bold = True
    p_maintitle.font.color.rgb = COLOR_WHITE
    p_maintitle.space_after = Pt(16)

    p_sub = tf1.add_paragraph()
    p_sub.text = (
        f"Dataset ID: {dataset_id}\n"
        f"Generated: {gen_time}   |   Architecture: Multi-Agent Autonomous Pipeline v1.0"
    )
    p_sub.font.size = Pt(11)
    p_sub.font.color.rgb = RGBColor(0xE2, 0xE8, 0xF0)

    # =========================================================================
    # SLIDE 2: EXECUTIVE SUMMARY
    # =========================================================================
    s2 = prs.slides.add_slide(blank_layout)
    _add_slide_header(s2, "Executive Summary & High-Level Metrics")
    _add_footer(s2, 2, TOTAL_SLIDES, file_name)

    # 5 KPI Cards: Rows, Columns, Quality, Domain, ML Status
    kpi_w = Inches(2.18)
    gap = Inches(0.2)
    top_pos = Inches(1.7)
    left_start = Inches(0.8)

    _add_kpi_card(s2, left_start, top_pos, kpi_w, Inches(1.4), f"{data['ingestion']['n_rows']:,}", "Ingested Records")
    _add_kpi_card(s2, left_start + (kpi_w + gap), top_pos, kpi_w, Inches(1.4), str(data["ingestion"]["n_columns"]), "Total Features")
    _add_kpi_card(s2, left_start + (kpi_w + gap) * 2, top_pos, kpi_w, Inches(1.4), f"{data['quality']['score']}/100", "Data Quality Score")
    _add_kpi_card(s2, left_start + (kpi_w + gap) * 3, top_pos, kpi_w, Inches(1.4), data["domain"]["name"], f"Domain ({int(data['domain']['confidence']*100)}% conf)")
    _add_kpi_card(s2, left_start + (kpi_w + gap) * 4, top_pos, kpi_w, Inches(1.4), data["ml"]["status"].upper(), "ML Pipeline Status")

    # Executive narrative takeaways card
    summary_bullets = [
        f"Dataset <b>{file_name}</b> successfully processed through the end-to-end multi-agent analytical pipeline.",
        f"Identified {len(data['columns'])} structured features classified into the {data['domain']['name']} domain.",
        f"Data quality evaluated at {data['quality']['score']}/100. " + (f"{len(data['quality']['issues'])} potential issues surfaced." if data['quality']['issues'] else "Zero structural defects found."),
    ]
    if data["pii_columns"]:
        summary_bullets.append(f"<b>Security Guard:</b> Detected {len(data['pii_columns'])} PII/identifier columns; all raw values strictly isolated from reporting outputs.")
    if data["insights"]:
        summary_bullets.append(f"Synthesized {len(data['insights'])} statistically verified business insights with numerical grounding.")

    _add_content_card(s2, Inches(0.8), Inches(3.4), Inches(11.733), Inches(3.2), "Key Operational Findings", summary_bullets)

    # =========================================================================
    # SLIDE 3: DATASET ARCHITECTURE & CATALOG
    # =========================================================================
    s3 = prs.slides.add_slide(blank_layout)
    _add_slide_header(s3, "Dataset Architecture & Feature Catalog")
    _add_footer(s3, 3, TOTAL_SLIDES, file_name)

    cat_headers = ["Feature Name", "Logical Type", "Null Count", "Null %", "Cardinality", "Privacy / Semantic Status"]
    cat_rows: list[list[str]] = []
    for col in data["columns"][:8]:  # Top 8 features
        cat_rows.append([
            col["name"],
            col["logical_type"],
            str(col["null_count"]),
            f"{col['null_pct']}%",
            str(col["unique_count"]),
            f"{col['semantic_label']} — {col['status_badge']}",
        ])

    cat_widths = [Inches(2.5), Inches(1.6), Inches(1.2), Inches(1.0), Inches(1.4), Inches(4.033)]
    _add_table(s3, Inches(0.8), Inches(1.7), Inches(11.733), Inches(4.2), cat_headers, cat_rows, cat_widths)

    privacy_note = [
        "Raw record values are strictly isolated in memory and never rendered into presentation or report deliverables.",
        f"PII Columns Protected: {', '.join(data['pii_columns']) if data['pii_columns'] else 'None detected'}.",
    ]
    _add_content_card(s3, Inches(0.8), Inches(5.95), Inches(11.733), Inches(0.85), "", privacy_note)

    # =========================================================================
    # SLIDE 4: DATA QUALITY & CLEANING OPERATIONS
    # =========================================================================
    s4 = prs.slides.add_slide(blank_layout)
    _add_slide_header(s4, "Data Quality Assessment & Cleaning Log")
    _add_footer(s4, 4, TOTAL_SLIDES, file_name)

    # Left: Quality Overview
    q_bullets = [
        f"Composite Quality Score: <b>{data['quality']['score']} / 100</b>",
    ]
    if data["quality"]["issues"]:
        for iss in data["quality"]["issues"]:
            q_bullets.append(f"Identified issue: {iss}")
    else:
        q_bullets.append("No critical missingness or schema integrity defects detected.")
    _add_content_card(s4, Inches(0.8), Inches(1.7), Inches(5.7), Inches(4.8), "Quality Evaluation", q_bullets)

    # Right: Cleaning Operations Executed
    c_bullets = []
    if data["cleaning"]["operations"]:
        for op in data["cleaning"]["operations"]:
            c_bullets.append(op)
    else:
        c_bullets.append("Dataset required zero destructive transformations; baseline data was clean.")
    c_bullets.append(f"Duplicate rows removed: {data['cleaning']['duplicates_removed']}")
    c_bullets.append(f"Total missing-value imputations: {data['cleaning']['imputation_count']}")
    _add_content_card(s4, Inches(6.8), Inches(1.7), Inches(5.733), Inches(4.8), "Cleaning Operations Applied", c_bullets)

    # =========================================================================
    # SLIDE 5: EXPLORATORY DATA ANALYSIS (EDA)
    # =========================================================================
    s5 = prs.slides.add_slide(blank_layout)
    _add_slide_header(s5, "Exploratory Data Analysis — Key Findings")
    _add_footer(s5, 5, TOTAL_SLIDES, file_name)

    # Summary Stats Table (Left)
    if data["eda"]["summary_stats"]:
        st_headers = ["Feature", "Mean", "Std", "Min", "Max"]
        st_rows = [[s["column"][:14], s["mean"], s["std"], s["min"], s["max"]] for s in data["eda"]["summary_stats"][:6]]
        st_widths = [Inches(1.8), Inches(1.0), Inches(1.0), Inches(1.0), Inches(1.0)]
        _add_table(s5, Inches(0.8), Inches(1.7), Inches(5.8), Inches(4.8), st_headers, st_rows, st_widths)
    else:
        _add_content_card(s5, Inches(0.8), Inches(1.7), Inches(5.8), Inches(4.8), "Summary Statistics", ["No numeric columns available for distribution summary."])

    # Top Correlations Table (Right)
    if data["eda"]["top_correlations"]:
        cr_headers = ["Feature A", "Feature B", "Correlation (r)"]
        cr_rows = [[c["feature_a"][:16], c["feature_b"][:16], f"{c['correlation']:+.3f}"] for c in data["eda"]["top_correlations"][:6]]
        cr_widths = [Inches(2.2), Inches(2.2), Inches(1.3)]
        _add_table(s5, Inches(6.8), Inches(1.7), Inches(5.7), Inches(4.8), cr_headers, cr_rows, cr_widths)
    else:
        _add_content_card(s5, Inches(6.8), Inches(1.7), Inches(5.7), Inches(4.8), "Feature Correlations", ["Zero or insufficient numeric feature pairs for correlation analysis."])

    # =========================================================================
    # SLIDE 6: KEY VISUALIZATIONS
    # =========================================================================
    s6 = prs.slides.add_slide(blank_layout)
    _add_slide_header(s6, "Visual Analytics & Distribution Highlights")
    _add_footer(s6, 6, TOTAL_SLIDES, file_name)

    charts = data["eda"]["chart_paths"]
    if len(charts) >= 2:
        # Two charts side by side
        cp1, cp2 = charts[0], charts[1]
        try:
            s6.shapes.add_picture(cp1, Inches(0.8), Inches(1.8), Inches(5.7), Inches(3.8))
            _add_content_card(s6, Inches(0.8), Inches(5.8), Inches(5.7), Inches(0.9), "", [f"Figure 1: {Path(cp1).stem.replace('_', ' ').title()}"])
        except Exception:
            _add_content_card(s6, Inches(0.8), Inches(1.8), Inches(5.7), Inches(3.8), "Visualization 1", [f"Chart available at: {Path(cp1).name}"])

        try:
            s6.shapes.add_picture(cp2, Inches(6.8), Inches(1.8), Inches(5.733), Inches(3.8))
            _add_content_card(s6, Inches(6.8), Inches(5.8), Inches(5.733), Inches(0.9), "", [f"Figure 2: {Path(cp2).stem.replace('_', ' ').title()}"])
        except Exception:
            _add_content_card(s6, Inches(6.8), Inches(1.8), Inches(5.733), Inches(3.8), "Visualization 2", [f"Chart available at: {Path(cp2).name}"])
    elif len(charts) == 1:
        cp = charts[0]
        try:
            s6.shapes.add_picture(cp, Inches(2.8), Inches(1.8), Inches(7.7), Inches(4.5))
            _add_content_card(s6, Inches(2.8), Inches(6.4), Inches(7.7), Inches(0.5), "", [f"Figure 1: {Path(cp).stem.replace('_', ' ').title()}"])
        except Exception:
            _add_content_card(s6, Inches(2.8), Inches(1.8), Inches(7.7), Inches(4.5), "Visualization 1", [f"Chart available at: {Path(cp).name}"])
    else:
        _add_content_card(s6, Inches(0.8), Inches(1.8), Inches(11.733), Inches(4.8), "Visualization Summary", ["No graphical charts were generated for this run due to data shape or configuration."])

    # =========================================================================
    # SLIDE 7: MACHINE LEARNING RESULTS
    # =========================================================================
    s7 = prs.slides.add_slide(blank_layout)
    _add_slide_header(s7, "Supervised Machine Learning Performance")
    _add_footer(s7, 7, TOTAL_SLIDES, file_name)

    ml_status = data["ml"]["status"]
    if ml_status == "trained":
        # Left: Model Overview & Metrics
        ml_bullets = [
            f"<b>Task Type:</b> {data['ml']['task_type']}",
            f"<b>Target Feature:</b> {data['ml']['target_column']}",
            f"<b>Selected Optimal Model:</b> {data['ml']['best_model']}",
        ]
        if data["ml"]["metrics"]:
            ml_bullets.append("<b>Evaluation Metrics:</b>")
            for m_k, m_v in data["ml"]["metrics"].items():
                ml_bullets.append(f"  • {m_k}: <b>{m_v}</b>")
        _add_content_card(s7, Inches(0.8), Inches(1.7), Inches(5.7), Inches(4.8), "Model Architecture & Performance", ml_bullets)

        # Right: Feature Importance
        if data["ml"]["feature_importances"]:
            fi_headers = ["Predictive Feature", "Relative Importance"]
            fi_rows = [[f[0], f"{f[1]:.4f}"] for f in data["ml"]["feature_importances"][:6]]
            fi_widths = [Inches(3.7), Inches(2.0)]
            _add_table(s7, Inches(6.8), Inches(1.7), Inches(5.733), Inches(4.8), fi_headers, fi_rows, fi_widths)
        else:
            _add_content_card(s7, Inches(6.8), Inches(1.7), Inches(5.733), Inches(4.8), "Feature Importance", ["No feature importance weights available."])
    else:
        # Informative notice card
        ml_notice_bullets = [
            f"<b>Pipeline Status:</b> {ml_status.upper()}",
            f"<b>Reason:</b> {data['ml']['reason'] or 'Analysis was not initiated for this dataset.'}",
            "Autonomous guardrails protect against invalid model training when sample size is insufficient, targets are absent, or data leakage risks exist.",
        ]
        _add_content_card(s7, Inches(0.8), Inches(1.7), Inches(11.733), Inches(4.8), "Machine Learning Safety Gate", ml_notice_bullets)

    # =========================================================================
    # SLIDE 8: KEY BUSINESS INSIGHTS
    # =========================================================================
    s8 = prs.slides.add_slide(blank_layout)
    _add_slide_header(s8, "Verified Business Insights")
    _add_footer(s8, 8, TOTAL_SLIDES, file_name)

    if data["insights"]:
        ins_cards = data["insights"][:3]  # Up to 3 featured insights
        num_c = len(ins_cards)
        if num_c == 1:
            card_w = Inches(11.733)
            card_gap = Inches(0.0)
        elif num_c == 2:
            card_gap = Inches(0.4)
            card_w = (Inches(11.733) - card_gap) / 2
        else:
            card_gap = Inches(0.24)
            card_w = Inches(3.75)

        for idx, ins in enumerate(ins_cards):
            left_pos = Inches(0.8) + (card_w + card_gap) * idx
            bullets = [
                f"<b>Category:</b> {ins['category'].upper()} ({int(ins['confidence']*100)}% conf)",
                ins["text"],
                f"<b>Evidence:</b> {ins['evidence']}",
            ]
            if ins["recommendation"]:
                bullets.append(f"<b>Recommendation:</b> {ins['recommendation']}")
            _add_content_card(s8, left_pos, Inches(1.7), card_w, Inches(4.8), f"Insight {ins['id']}", bullets)
    else:
        _add_content_card(s8, Inches(0.8), Inches(1.7), Inches(11.733), Inches(4.8), "Business Insights", ["Limited statistical evidence was available to derive high-confidence business insights."])

    # =========================================================================
    # SLIDE 9: ACTIONABLE RECOMMENDATIONS
    # =========================================================================
    s9 = prs.slides.add_slide(blank_layout)
    _add_slide_header(s9, "Strategic & Actionable Recommendations")
    _add_footer(s9, 9, TOTAL_SLIDES, file_name)

    rec_bullets: list[str] = []
    if data["recommendations"]:
        for idx, r in enumerate(data["recommendations"][:5], start=1):
            rec_bullets.append(f"<b>Priority {idx}:</b> {r}")
    else:
        rec_bullets.append("Maintain data collection consistency and re-evaluate recommendations upon expanded sample volume.")
    _add_content_card(s9, Inches(0.8), Inches(1.7), Inches(11.733), Inches(4.8), "Strategic Action Plan", rec_bullets)

    # =========================================================================
    # SLIDE 10: ANALYTICAL LIMITATIONS & GUARDRAILS
    # =========================================================================
    s10 = prs.slides.add_slide(blank_layout)
    _add_slide_header(s10, "Analytical Limitations & Data Warnings")
    _add_footer(s10, 10, TOTAL_SLIDES, file_name)

    lim_bullets: list[str] = []
    if data["limitations"]:
        for l in data["limitations"][:6]:
            lim_bullets.append(l)
    else:
        lim_bullets.append("No material data caveats or pipeline anomalies observed during analysis.")
    _add_content_card(s10, Inches(0.8), Inches(1.7), Inches(11.733), Inches(4.8), "Observed Caveats & Constraints", lim_bullets)

    # =========================================================================
    # SLIDE 11: APPENDIX & PROVENANCE
    # =========================================================================
    s11 = prs.slides.add_slide(blank_layout)
    _add_slide_header(s11, "Appendix — Technical Provenance & Telemetry")
    _add_footer(s11, 11, TOTAL_SLIDES, file_name)

    app_headers = ["Technical Property", "Runtime Value"]
    app_rows = [
        ["Run / Dataset ID", data["metadata"]["dataset_id"]],
        ["Source File Name", data["metadata"]["file_name"]],
        ["Dataset Hash (SHA-256)", data["metadata"]["dataset_hash"] or "Not Calculated"],
        ["DIO Schema Version", data["metadata"]["schema_version"]],
        ["Generation Timestamp", data["metadata"]["generated_at"]],
        ["Report Status", "Verified Deliverables Generated"],
    ]
    app_widths = [Inches(3.5), Inches(8.233)]
    _add_table(s11, Inches(0.8), Inches(1.7), Inches(11.733), Inches(4.8), app_headers, app_rows, app_widths)

    prs.save(str(dest_path))
    return dest_path
