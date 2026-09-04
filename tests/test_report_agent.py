"""
tests/test_report_agent.py
==========================
Unit tests for Phase 8: Report Agent.
Verifies PDF and PPTX deliverable generation, DIO contracts, file integrity,
missing sections handling, and generator resilience.
"""

from __future__ import annotations

import re
from pathlib import Path
from unittest.mock import patch
import pytest
from pptx import Presentation

from agents.report.report_agent import ReportAgent
from agents.report.pdf_generator import generate_pdf_report
from agents.report.pptx_generator import generate_pptx_report
from core.dio import DIO


def _create_test_dio(file_name: str = "retail_test.csv") -> DIO:
    """Helper to construct a comprehensive synthetic DIO for unit testing."""
    dio = DIO.create_empty(file_name, dataset_hash="abc123hash")
    dio.ingestion = {
        "n_rows": 250,
        "n_columns": 5,
        "file_type": "csv",
        "encoding": "utf-8",
    }
    dio.domain_guess = {"domain": "retail", "confidence": 0.89}
    dio.quality = {"score": 94, "issues": ["1 duplicate row detected"]}
    dio.columns = [
        {"name": "order_id", "logical_type": "integer", "null_count": 0, "null_pct": 0.0, "unique_count": 250, "is_identifier": True, "is_pii": False, "semantic_label": "id"},
        {"name": "customer_email", "logical_type": "string", "null_count": 0, "null_pct": 0.0, "unique_count": 240, "is_identifier": False, "is_pii": True, "pii_type": "email", "semantic_label": "email_address"},
        {"name": "revenue", "logical_type": "float", "null_count": 2, "null_pct": 0.8, "unique_count": 210, "is_identifier": False, "is_pii": False, "semantic_label": "monetary_amount"},
        {"name": "category", "logical_type": "category", "null_count": 0, "null_pct": 0.0, "unique_count": 4, "is_identifier": False, "is_pii": False, "semantic_label": "category"},
        {"name": "order_date", "logical_type": "date", "null_count": 0, "null_pct": 0.0, "unique_count": 180, "is_identifier": False, "is_pii": False, "semantic_label": "date"},
    ]
    dio.cleaning_log = [
        {"action": "drop_duplicates", "count": 1, "details": "Removed 1 duplicate row"},
        {"action": "impute_missing", "column": "revenue", "method": "median", "replacement_value": 45.50},
    ]
    dio.eda = {
        "summary_stats": {
            "revenue": {"mean": 85.20, "std": 22.40, "min": 15.0, "median": 82.0, "max": 240.0},
        },
        "correlations": {
            "top_correlations": [
                {"feature_a": "revenue", "feature_b": "order_id", "pearson_correlation": 0.12},
            ],
        },
        "charts": [],
    }
    dio.ml = {
        "status": "trained",
        "task_type": "regression",
        "target_column": "revenue",
        "best_model": "RandomForestRegressor",
        "metrics": {"r2_score": 0.7845, "rmse": 10.25},
        "feature_importance": {"category": 0.65, "order_id": 0.35},
        "reason": "Successfully trained candidate models",
    }
    dio.insights = [
        {
            "id": "ins_001",
            "category": "distribution",
            "text": "Average revenue was 85.20 with median at 82.00.",
            "confidence": 0.94,
            "evidence": "eda.summary_stats.revenue.mean",
            "recommendation": "Target high-margin retail product lines.",
            "grounded_numbers": [85.20, 82.00],
        },
    ]
    return dio


def test_basic_report_generation(tmp_path: Path):
    """Verify that ReportAgent generates both PDF and PPTX deliverables."""
    dio = _create_test_dio()
    agent = ReportAgent()

    _, dio_out = agent.run(None, dio, run_dir=tmp_path)

    assert dio_out["progress"]["report"] == "FINISHED"
    assert dio_out["reports"]["status"] == "completed"

    pdf_path = Path(dio_out["reports"]["pdf_path"])
    pptx_path = Path(dio_out["reports"]["pptx_path"])

    assert pdf_path.is_file()
    assert pptx_path.is_file()
    assert pdf_path.stat().st_size > 0
    assert pptx_path.stat().st_size > 0


def test_file_integrity_and_readability(tmp_path: Path):
    """Verify that generated PDF and PPTX files can be cleanly opened and inspected."""
    dio = _create_test_dio()
    agent = ReportAgent()

    _, dio_out = agent.run(None, dio, run_dir=tmp_path)
    pdf_path = Path(dio_out["reports"]["pdf_path"])
    pptx_path = Path(dio_out["reports"]["pptx_path"])

    # 1. PDF integrity: check magic bytes and page count regex
    pdf_bytes = pdf_path.read_bytes()
    assert pdf_bytes.startswith(b"%PDF-")
    assert b"%%EOF" in pdf_bytes
    page_count = len(re.findall(rb"/Type\s*/Page\b", pdf_bytes))
    assert page_count >= 1

    # 2. PPTX integrity: open with python-pptx and verify slide count
    prs = Presentation(pptx_path)
    assert len(prs.slides) == 11
    # Check that title slide contains dataset name
    title_text = ""
    for shape in prs.slides[0].shapes:
        if shape.has_text_frame:
            title_text += shape.text_frame.text
    assert "retail_test.csv" in title_text


def test_dio_reports_schema_contract(tmp_path: Path):
    """Verify that dio['reports'] and dio['artifacts'] match the exact expected schema."""
    dio = _create_test_dio()
    agent = ReportAgent()

    _, dio_out = agent.run(None, dio, run_dir=tmp_path)

    reports = dio_out["reports"]
    assert "status" in reports
    assert "pdf_path" in reports
    assert "pptx_path" in reports
    assert "generated_at" in reports
    assert "warnings" in reports
    assert reports["status"] == "completed"
    assert isinstance(reports["warnings"], list)

    assert dio_out["artifacts"]["pdf_report"] == reports["pdf_path"]
    assert dio_out["artifacts"]["pptx_report"] == reports["pptx_path"]

    # Check metrics & decision log
    assert any(m["agent"] == "report" for m in dio_out["agent_metrics"])
    assert any(d["agent"] == "report" for d in dio_out["decision_log"])


def test_missing_ml_handling(tmp_path: Path):
    """Verify that ReportAgent gracefully handles missing or insufficient ML sections."""
    dio = _create_test_dio()
    # Simulate insufficient data for ML
    dio.ml = {
        "status": "insufficient_data",
        "task_type": "none",
        "target_column": None,
        "reason": "Dataset rows (25) below minimum threshold (30)",
    }

    agent = ReportAgent()
    _, dio_out = agent.run(None, dio, run_dir=tmp_path)

    assert dio_out["reports"]["status"] == "completed"
    pdf_bytes = Path(dio_out["reports"]["pdf_path"]).read_bytes()
    assert b"insufficient" in pdf_bytes.lower() or b"sample size" in pdf_bytes.lower()

    # PPTX check
    prs = Presentation(dio_out["reports"]["pptx_path"])
    ml_slide_text = "".join(
        shape.text_frame.text for shape in prs.slides[6].shapes if shape.has_text_frame
    )
    assert "INSUFFICIENT_DATA" in ml_slide_text or "insufficient" in ml_slide_text.lower()


def test_missing_charts_handling(tmp_path: Path):
    """Verify that report generation succeeds with an empty chart list or missing chart files."""
    dio = _create_test_dio()
    dio.artifacts["chart_paths"] = ["/non/existent/path/chart_1.png", "/fake/chart_2.png"]
    dio.eda["charts"] = ["/non/existent/path/chart_1.png"]

    agent = ReportAgent()
    _, dio_out = agent.run(None, dio, run_dir=tmp_path)

    assert dio_out["reports"]["status"] == "completed"
    assert Path(dio_out["reports"]["pdf_path"]).is_file()
    assert Path(dio_out["reports"]["pptx_path"]).is_file()


def test_missing_insights_handling(tmp_path: Path):
    """Verify that ReportAgent succeeds with zero insights without crashing."""
    dio = _create_test_dio()
    dio.insights = []

    agent = ReportAgent()
    _, dio_out = agent.run(None, dio, run_dir=tmp_path)

    assert dio_out["reports"]["status"] == "completed"
    pdf_bytes = Path(dio_out["reports"]["pdf_path"]).read_bytes()
    assert b"verified insights" in pdf_bytes.lower() or b"insights" in pdf_bytes.lower()


def test_minimal_empty_dio_handling(tmp_path: Path):
    """Verify that a minimal DIO with missing optional sections generates valid reports."""
    dio = DIO.create_empty("minimal_dataset.csv")
    agent = ReportAgent()

    _, dio_out = agent.run(None, dio, run_dir=tmp_path)

    assert dio_out["reports"]["status"] == "completed"
    assert Path(dio_out["reports"]["pdf_path"]).is_file()
    assert Path(dio_out["reports"]["pptx_path"]).is_file()


def test_independent_generator_resilience_pdf_failure(tmp_path: Path):
    """Verify that if PDF generation fails, PPTX generation is preserved and partial status is recorded."""
    dio = _create_test_dio()
    agent = ReportAgent()

    with patch("agents.report.report_agent.generate_pdf_report", side_effect=RuntimeError("Simulated PDF crash")):
        _, dio_out = agent.run(None, dio, run_dir=tmp_path)

    reports = dio_out["reports"]
    assert reports["status"] == "partial"
    assert reports["pdf_path"] is None
    assert reports["pptx_path"] is not None
    assert Path(reports["pptx_path"]).is_file()
    assert any(e.get("code") == "RPT_001" for e in dio_out["errors"])


def test_independent_generator_resilience_pptx_failure(tmp_path: Path):
    """Verify that if PPTX generation fails, PDF generation is preserved and partial status is recorded."""
    dio = _create_test_dio()
    agent = ReportAgent()

    with patch("agents.report.report_agent.generate_pptx_report", side_effect=RuntimeError("Simulated PPTX crash")):
        _, dio_out = agent.run(None, dio, run_dir=tmp_path)

    reports = dio_out["reports"]
    assert reports["status"] == "partial"
    assert reports["pdf_path"] is not None
    assert Path(reports["pdf_path"]).is_file()
    assert reports["pptx_path"] is None
    assert any(e.get("code") == "RPT_002" for e in dio_out["errors"])
