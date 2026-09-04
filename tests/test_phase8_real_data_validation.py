"""
tests/test_phase8_real_data_validation.py
=========================================
Real-data validation suite for Phase 8: Report Agent.
Executes the full end-to-end multi-agent pipeline:
  DataRouter -> IntelligenceAgent -> CleaningAgent -> EDAAgent -> MLAgent -> InsightAgent -> ReportAgent
across all 5 project benchmark datasets and verifies:
1. PDF and PPTX artifact creation & readability
2. Zero PII leakage in deliverables
3. Strict upstream DIO immutability
4. Faithful rendering of verified insights & ML status
"""

from __future__ import annotations

import copy
from pathlib import Path
import pytest
from pptx import Presentation

from core.data_router import DataRouter
from agents.intelligence.intelligence_agent import IntelligenceAgent
from agents.cleaning.cleaning_agent import CleaningAgent
from agents.eda.eda_agent import EDAAgent
from agents.ml.ml_agent import MLAgent
from agents.insight.insight_agent import InsightAgent
from agents.report.report_agent import ReportAgent

DATASETS = [
    "retail_sales.csv",
    "healthcare_patients.csv",
    "financial_loans.csv",
    "mixed_messy_data.csv",
    "ambiguous_dates_pii.csv",
]


def _extract_all_pptx_text(pptx_path: Path | str) -> str:
    """Helper to extract all text from all shapes and tables in a PPTX."""
    prs = Presentation(pptx_path)
    texts: list[str] = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                for p in shape.text_frame.paragraphs:
                    texts.append(p.text)
            elif shape.has_table:
                for row in shape.table.rows:
                    for cell in row.cells:
                        texts.append(cell.text)
    return " ".join(texts)


@pytest.mark.parametrize("dataset_name", DATASETS)
def test_phase8_real_data_pipeline_validation(dataset_name: str, tmp_path: Path):
    project_root = Path(__file__).resolve().parent.parent
    csv_path = project_root / "data" / "sample" / dataset_name
    assert csv_path.exists(), f"Sample dataset not found: {csv_path}"

    # 1. Ingestion
    router = DataRouter()
    df, dio, run_dir = router.ingest(csv_path, base_runs_dir=tmp_path)

    # 2. Intelligence Agent
    intel_agent = IntelligenceAgent()
    df, dio = intel_agent.run(df, dio)

    # 3. Cleaning Agent
    clean_agent = CleaningAgent()
    cleaned_df, dio = clean_agent.run(df, dio, run_dir=run_dir)

    # 4. EDA Agent
    eda_agent = EDAAgent()
    cleaned_df, dio = eda_agent.run(cleaned_df, dio, run_dir=run_dir)

    # 5. ML Agent
    ml_agent = MLAgent()
    cleaned_df, dio = ml_agent.run(cleaned_df, dio, run_dir=run_dir)

    # 6. Insight Agent
    insight_agent = InsightAgent()
    cleaned_df, dio = insight_agent.run(cleaned_df, dio, run_dir=run_dir)

    # Snapshot upstream DIO before Report Agent execution
    upstream_keys = [
        "schema_version",
        "dataset_id",
        "dataset_hash",
        "file_name",
        "ingestion",
        "columns",
        "date_columns",
        "domain_guess",
        "quality",
        "cleaning_log",
        "eda",
        "ml",
        "insights",
    ]
    upstream_snapshot = {k: copy.deepcopy(dio[k]) for k in upstream_keys}

    # 7. Report Agent (Phase 8)
    report_agent = ReportAgent()
    _, dio = report_agent.run(cleaned_df, dio, run_dir=run_dir)

    # -------------------------------------------------------------------------
    # VERIFICATIONS
    # -------------------------------------------------------------------------
    # A. Upstream DIO Immutability
    for k in upstream_keys:
        assert dio[k] == upstream_snapshot[k], f"ReportAgent mutated upstream DIO section '{k}' on {dataset_name}!"

    # B. Report Section Contract
    assert dio["progress"]["report"] == "FINISHED"
    assert dio["reports"]["status"] == "completed"
    assert dio["reports"]["pdf_path"] is not None
    assert dio["reports"]["pptx_path"] is not None

    pdf_path = Path(dio["reports"]["pdf_path"])
    pptx_path = Path(dio["reports"]["pptx_path"])

    # C. File Existence & Integrity
    assert pdf_path.is_file(), f"PDF report file does not exist: {pdf_path}"
    assert pptx_path.is_file(), f"PPTX report file does not exist: {pptx_path}"
    assert pdf_path.stat().st_size > 0, "PDF file is empty"
    assert pptx_path.stat().st_size > 0, "PPTX file is empty"

    # D. PDF Readability
    pdf_bytes = pdf_path.read_bytes()
    assert pdf_bytes.startswith(b"%PDF-"), "Invalid PDF header magic bytes"
    assert b"%%EOF" in pdf_bytes, "PDF missing EOF terminator"

    # E. PPTX Readability
    prs = Presentation(pptx_path)
    assert len(prs.slides) == 11, f"Expected 11 slides in presentation, got {len(prs.slides)}"

    # F. PII Protection Check
    # Collect raw PII values from the original raw DataFrame to ensure none leaked
    pii_columns = [c["name"] for c in dio["columns"] if c.get("is_pii") or c.get("is_identifier")]
    pdf_text = pdf_bytes.decode("latin1", errors="ignore")
    pptx_text = _extract_all_pptx_text(pptx_path)

    for col_name in pii_columns:
        if col_name in df.columns:
            for raw_val in df[col_name].dropna().unique()[:5]:
                val_str = str(raw_val).strip()
                # Skip trivial strings or short numbers that could coincidentally match dates/indices
                if len(val_str) > 4 and not val_str.isdigit():
                    assert val_str not in pdf_text, f"PII value '{val_str}' leaked into PDF report for {dataset_name}!"
                    assert val_str not in pptx_text, f"PII value '{val_str}' leaked into PPTX report for {dataset_name}!"

    # G. Insights Rendering Check
    if dio["insights"]:
        first_insight = dio["insights"][0]
        if isinstance(first_insight, dict):
            ins_id = first_insight.get("id", "")
            ins_text_snippet = first_insight.get("text", "")[:25]
            assert (
                ins_id in pdf_text
                or ins_id in pptx_text
                or ins_text_snippet in pdf_text
                or ins_text_snippet in pptx_text
            ), f"Insight '{ins_id}' was not rendered in deliverables for {dataset_name}!"

    # H. ML Status Representation Check
    ml_status = dio["ml"]["status"]
    assert ml_status.upper() in pptx_text or ml_status.lower() in pdf_text.lower()
