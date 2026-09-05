"""
tests/benchmark/evaluators/report_evaluator.py
=============================================
Evaluator for Metric 9: Report Fidelity Evaluation.
Verifies:
  - report deliverables exist and have non-empty size (> 1000 bytes)
  - objective pipeline facts (row count, column count, quality score, domain)
    match extracted text from PDF and PPTX artifacts
  - zero raw PII strings leaked in deliverable documents
Produces:
  - report_fidelity_accuracy
  - pdf_generated
  - pptx_generated
  - facts_verified_count
  - facts_total_count
  - factual_mismatches
"""

from __future__ import annotations

from dataclasses import asdict, dataclass, field
from pathlib import Path
import re
from typing import Any
from pptx import Presentation

from core.dio import DIO
from tests.benchmark.ground_truth import DatasetGroundTruth

# PDF binary text stream extractor
PDF_TEXT_RE = re.compile(rb"\(((?:\\.|[^\\()])*)\)")


def extract_pdf_text(pdf_path: Path | str) -> str:
    """Extract decoded text from uncompressed text objects in PDF."""
    path_obj = Path(pdf_path)
    if not path_obj.exists():
        return ""
    data = path_obj.read_bytes()
    matches = PDF_TEXT_RE.findall(data)
    decoded_parts: list[str] = []
    for m in matches:
        try:
            # Handle standard escaped characters \( \) \\
            cleaned = m.replace(b"\\(", b"(").replace(b"\\)", b")").replace(b"\\\\", b"\\")
            decoded_parts.append(cleaned.decode("latin-1", errors="ignore"))
        except Exception:
            continue
    return " ".join(decoded_parts)


def extract_pptx_text(pptx_path: Path | str) -> str:
    """Extract all text from shapes and tables in a PPTX."""
    path_obj = Path(pptx_path)
    if not path_obj.exists():
        return ""
    prs = Presentation(str(path_obj))
    texts: list[str] = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                for p in shape.text_frame.paragraphs:
                    texts.append(p.text)
            elif shape.has_table:
                for row in shape.table.rows:
                    for cell in row.cells:
                        texts.append(cell.text)
    return " ".join(texts)


@dataclass
class ReportEvaluationResult:
    report_fidelity_accuracy: float
    pdf_generated: bool
    pptx_generated: bool
    facts_verified_count: int
    facts_total_count: int
    factual_mismatches: list[dict[str, Any]] = field(default_factory=list)

    def to_dict(self) -> dict[str, Any]:
        return asdict(self)


class ReportEvaluator:
    """Evaluates artifact persistence and factual fidelity of generated reports."""

    def evaluate(
        self,
        dio: DIO,
        ground_truth: DatasetGroundTruth,
        run_dir: Path | None = None,
    ) -> ReportEvaluationResult:
        artifacts = dio.get("artifacts", {})
        pdf_path = artifacts.get("pdf_report")
        pptx_path = artifacts.get("pptx_report")

        pdf_exists = bool(pdf_path and Path(pdf_path).exists() and Path(pdf_path).stat().st_size > 1000)
        pptx_exists = bool(pptx_path and Path(pptx_path).exists() and Path(pptx_path).stat().st_size > 1000)

        facts_verified = 0
        facts_total = 0
        mismatches: list[dict[str, Any]] = []

        # Artifact existence checks
        facts_total += 2
        if pdf_exists:
            facts_verified += 1
        else:
            mismatches.append({"artifact": "pdf_report", "error": "PDF report missing or size <= 1000 bytes"})

        if pptx_exists:
            facts_verified += 1
        else:
            mismatches.append({"artifact": "pptx_report", "error": "PPTX report missing or size <= 1000 bytes"})

        # If reports exist, extract text and check objective facts
        extracted_text = ""
        if pdf_exists and pdf_path:
            extracted_text += " " + extract_pdf_text(pdf_path)
        if pptx_exists and pptx_path:
            extracted_text += " " + extract_pptx_text(pptx_path)

        if extracted_text.strip():
            # Fact 1: Row count
            n_rows = dio.get("ingestion", {}).get("n_rows")
            if n_rows is not None:
                facts_total += 1
                if str(n_rows) in extracted_text:
                    facts_verified += 1
                else:
                    mismatches.append({"fact": "n_rows", "value": n_rows, "error": f"Row count {n_rows} not found in report text"})

            # Fact 2: Column count
            n_cols = dio.get("ingestion", {}).get("n_columns")
            if n_cols is not None:
                facts_total += 1
                if str(n_cols) in extracted_text:
                    facts_verified += 1
                else:
                    mismatches.append({"fact": "n_columns", "value": n_cols, "error": f"Column count {n_cols} not found in report text"})

            # Fact 3: Quality score
            q_score = dio.get("quality", {}).get("score")
            if q_score is not None:
                facts_total += 1
                if str(q_score) in extracted_text:
                    facts_verified += 1
                else:
                    mismatches.append({"fact": "quality_score", "value": q_score, "error": f"Quality score {q_score} not found in report text"})

            # Fact 4: Domain
            domain = dio.get("domain_guess", {}).get("domain", "")
            if domain:
                facts_total += 1
                if domain.lower() in extracted_text.lower():
                    facts_verified += 1
                else:
                    mismatches.append({"fact": "domain", "value": domain, "error": f"Domain '{domain}' not found in report text"})

            # Fact 5: Zero Raw PII in report text
            for sensitive_val in ground_truth.sensitive_raw_values:
                if sensitive_val in extracted_text:
                    facts_total += 1
                    mismatches.append({
                        "security_violation": "Raw PII leakage in generated report text!",
                        "error": "Sensitive value appeared in report artifact text",
                    })

        accuracy = round(facts_verified / facts_total, 4) if facts_total > 0 else 1.0

        return ReportEvaluationResult(
            report_fidelity_accuracy=accuracy,
            pdf_generated=pdf_exists,
            pptx_generated=pptx_exists,
            facts_verified_count=facts_verified,
            facts_total_count=facts_total,
            factual_mismatches=mismatches,
        )
